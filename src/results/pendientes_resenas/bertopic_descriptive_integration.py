"""Fit the descriptive BERTopic model and integrate observed review labels."""

from __future__ import annotations

import ast
import argparse
import json
import time
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.preprocessing import normalize

from bertopic_validation_grid import diversity, make_model, topic_npmi, top_words_from_model


ROOT = Path(__file__).resolve().parents[2]
PROCESSED = ROOT / "src" / "data" / "processed"
OUTPUT = ROOT / "results" / "pendientes_resenas" / "descriptive_cap_50"
EMBEDDING_MODEL = "all-MiniLM-L6-v2"
RANDOM_STATE = 42
EMBEDDING_BATCH_SIZE = 16
MIN_CLUSTER_SIZES = (50, 100, 200)


def load_catalog() -> pd.DataFrame:
    catalog = pd.read_csv(
        PROCESSED / "games_fixed_clean.csv",
        usecols=["appid", "genres"],
        dtype={"appid": "string", "genres": "string"},
        keep_default_na=False,
    ).drop_duplicates("appid", keep="first")

    def primary_genre(value: str) -> str:
        if not value:
            return "Unknown"
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            try:
                parsed = ast.literal_eval(value)
            except (ValueError, SyntaxError):
                parsed = [part.strip() for part in value.split(",") if part.strip()]
        if isinstance(parsed, list):
            return str(parsed[0]).strip() if parsed else "Unknown"
        return str(parsed).strip() if parsed else "Unknown"

    catalog["primary_genre"] = catalog["genres"].map(primary_genre)
    return catalog[["appid", "primary_genre"]]


def summarize(documents: list[str], labels: np.ndarray,
              topic_words: dict[int, list[str]]) -> dict:
    labels = np.asarray(labels)
    sizes = {
        str(int(topic)): int((labels == topic).sum())
        for topic in sorted(set(labels.tolist()))
        if topic != -1
    }
    return {
        "topic_count": len(topic_words),
        "npmi": topic_npmi(documents, topic_words),
        "topic_diversity": diversity(topic_words),
        "outlier_fraction": float(np.mean(labels == -1)),
        "topic_sizes": sizes,
        "largest_topic_share_non_outlier": (
            max(sizes.values()) / sum(sizes.values()) if sizes else 1.0
        ),
    }


def topic_genre_integration(
    corpus: pd.DataFrame,
    labels: np.ndarray,
    topic_words: dict[int, list[str]],
) -> pd.DataFrame:
    data = corpus[["appid", "voted_up", "primary_genre"]].copy()
    data["topic"] = np.asarray(labels, dtype=int)
    data["negative"] = data["voted_up"].eq(0).astype(int)
    data = data.loc[data["primary_genre"].ne("Unknown")].copy()
    segment_totals = data.groupby("primary_genre").agg(
        N_s=("negative", "size"), negative_s=("negative", "sum")
    )
    eligible = segment_totals.loc[segment_totals["N_s"].ge(100)]
    z = 1.959963984540054
    rows = []

    for genre, segment in eligible.iterrows():
        group = data.loc[data["primary_genre"].eq(genre)]
        n_s = len(group)
        q_s = float(group["negative"].mean())
        topic_group = group.loc[group["topic"].ne(-1)]
        phi = topic_group["topic"].value_counts() / n_s
        top_three = set(phi.sort_values(ascending=False).head(3).index)

        for topic_id, topic_frame in topic_group.groupby("topic"):
            n_ks = len(topic_frame)
            if n_ks < 10:
                continue
            negative_count = int(topic_frame["negative"].sum())
            q_ks = negative_count / n_ks
            phi_ks = n_ks / n_s
            denominator = 1 + z * z / n_ks
            center = (q_ks + z * z / (2 * n_ks)) / denominator
            margin = (
                z
                * np.sqrt(q_ks * (1 - q_ks) / n_ks + z * z / (4 * n_ks**2))
                / denominator
            )
            game_rates = topic_frame.groupby("appid").agg(
                mentions=("negative", "size"), negative_count=("negative", "sum")
            )
            eligible_games = game_rates.loc[game_rates["mentions"].ge(5)]
            q_per_game = (
                float((eligible_games["negative_count"] / eligible_games["mentions"]).mean())
                if not eligible_games.empty
                else np.nan
            )
            lower, upper = max(0.0, center - margin), min(1.0, center + margin)
            top3 = topic_id in top_three
            rows.append(
                {
                    "topic": int(topic_id),
                    "primary_genre": genre,
                    "topic_words": " | ".join(topic_words.get(int(topic_id), [])),
                    "n_k_s": int(n_ks),
                    "negative_count": negative_count,
                    "N_s": int(n_s),
                    "phi_k_s": float(phi_ks),
                    "q_k_s": float(q_ks),
                    "q_s": q_s,
                    "I_k_s": float(phi_ks * q_ks),
                    "q_tilde_m20": float((n_ks * q_ks + 20 * q_s) / (n_ks + 20)),
                    "wilson_lower_95": lower,
                    "wilson_upper_95": upper,
                    "top_3_phi": top3,
                    "unmet_need": bool(lower > q_s and top3),
                    "valued_attribute": bool(upper < q_s and top3),
                    "games_with_at_least_5_mentions": int(len(eligible_games)),
                    "q_mean_per_game_min_5_mentions": q_per_game,
                }
            )
    return pd.DataFrame(rows)


def representative_topics(
    corpus: pd.DataFrame,
    labels: np.ndarray,
    embeddings: np.ndarray,
    topic_words: dict[int, list[str]],
) -> pd.DataFrame:
    labels = np.asarray(labels)
    normalized = normalize(embeddings)
    rows = []
    for topic_id, words in topic_words.items():
        indices = np.flatnonzero(labels == topic_id)
        if not len(indices):
            continue
        centroid = normalized[indices].mean(axis=0)
        similarities = normalized[indices] @ centroid
        chosen = indices[np.argsort(similarities)[-3:][::-1]]
        snippets = [
            " ".join(str(corpus.iloc[index]["review"]).split())[:300]
            for index in chosen
        ]
        rows.append(
            {
                "topic": int(topic_id),
                "size": int(len(indices)),
                "top_words": " | ".join(words),
                "representative_review_1": snippets[0],
                "representative_review_2": snippets[1] if len(snippets) > 1 else "",
                "representative_review_3": snippets[2] if len(snippets) > 2 else "",
            }
        )
    return pd.DataFrame(rows).sort_values("topic").reset_index(drop=True)


def write_review_workbook(topic_rows: pd.DataFrame, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    review_sheet = workbook.active
    review_sheet.title = "Revisión"
    headers = [
        "Id del tema",
        "Tamaño",
        "10 palabras principales",
        "Reseña representativa 1",
        "Reseña representativa 2",
        "Reseña representativa 3",
        "Revisor 1",
        "Revisor 2",
        "Revisor 3",
        "Nombre propuesto del tema",
    ]
    review_sheet.append(headers)
    for row in topic_rows.itertuples(index=False):
        review_sheet.append(
            [
                int(row.topic),
                int(row.size),
                row.top_words,
                row.representative_review_1,
                row.representative_review_2,
                row.representative_review_3,
                None,
                None,
                None,
                None,
            ]
        )

    last_review_row = review_sheet.max_row
    validation = DataValidation(
        type="list",
        formula1='"Interpretable,Parcialmente,No interpretable"',
        allow_blank=True,
    )
    review_sheet.add_data_validation(validation)
    validation.add(f"G2:I{last_review_row}")
    review_sheet.freeze_panes = "A2"
    review_sheet.auto_filter.ref = f"A1:J{last_review_row}"
    widths = [13, 12, 72, 54, 54, 54, 20, 20, 20, 34]
    for column, width in enumerate(widths, start=1):
        review_sheet.column_dimensions[
            review_sheet.cell(row=1, column=column).column_letter
        ].width = width
    for row in review_sheet.iter_rows(min_row=2, min_col=3, max_col=6):
        for cell in row:
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")

    agreement = workbook.create_sheet("Acuerdo")
    agreement["A1"] = "Resumen de etiquetas por revisor"
    agreement.append(["Etiqueta", "Revisor 1", "Revisor 2", "Revisor 3"])
    labels = ["Interpretable", "Parcialmente", "No interpretable"]
    for row_number, label in enumerate(labels, start=3):
        agreement.cell(row=row_number, column=1, value=label)
        for reviewer_col, source_col in zip(("B", "C", "D"), ("G", "H", "I")):
            agreement[f"{reviewer_col}{row_number}"] = (
                f'=IFERROR(COUNTIF(\'Revisión\'!${source_col}$2:${source_col}${last_review_row},'
                f'$A{row_number})/COUNTIF(\'Revisión\'!${source_col}$2:${source_col}${last_review_row},'
                f'"<>"),0)'
            )
    agreement["A7"] = "Porcentaje de temas con acuerdo total (solo temas con 3 etiquetas)"
    agreement["B7"] = (
        f"=IFERROR(SUMPRODUCT(--('Revisión'!G2:G{last_review_row}="
        f"'Revisión'!H2:H{last_review_row}),--('Revisión'!H2:H{last_review_row}="
        f"'Revisión'!I2:I{last_review_row}),--('Revisión'!G2:G{last_review_row}<>\"\"))/"
        f"COUNTIFS('Revisión'!G2:G{last_review_row},\"<>\","
        f"'Revisión'!H2:H{last_review_row},\"<>\",'Revisión'!I2:I{last_review_row},\"<>\"),0)"
    )
    agreement["A8"] = "Kappa de Fleiss (temas con 3 etiquetas)"
    agreement["B8"] = "=IFERROR((B9-E10)/(1-E10),0)"
    agreement["A9"] = "Acuerdo observado medio (P-barra)"
    agreement["A10"] = "Acuerdo esperado (P-e)"
    agreement["A12"] = "Cálculo por tema (etiquetar las tres columnas de cada fila)"
    agreement.append(["Id del tema", "Interpretable", "Parcialmente", "No interpretable", "P_i"])
    first_helper_row = 14
    for offset, review_row in enumerate(range(2, last_review_row + 1)):
        helper = first_helper_row + offset
        agreement.cell(
            helper, 1, value=f"='Revisión'!A{review_row}"
        )
        for col, source_label in zip(("B", "C", "D"), labels):
            agreement[f"{col}{helper}"] = (
                f'=IF(COUNTA(\'Revisión\'!G{review_row}:I{review_row})=3,'
                f'COUNTIF(\'Revisión\'!G{review_row}:I{review_row},"{source_label}"),0)'
            )
        agreement[f"E{helper}"] = (
            f'=IF(SUM(B{helper}:D{helper})=3,'
            f'(SUMSQ(B{helper}:D{helper})-3)/(3*2),"")'
        )
    last_helper_row = first_helper_row + len(topic_rows) - 1
    agreement["B9"] = f"=IFERROR(AVERAGE(E{first_helper_row}:E{last_helper_row}),0)"
    complete_topics = (
        f'COUNTIFS(Revisión!G2:G{last_review_row},"<>",'
        f'Revisión!H2:H{last_review_row},"<>",Revisión!I2:I{last_review_row},"<>")'
    )
    for cell, category_col in zip(("B10", "C10", "D10"), ("B", "C", "D")):
        agreement[cell] = (
            f"=IFERROR((SUM({category_col}{first_helper_row}:"
            f"{category_col}{last_helper_row})/(3*{complete_topics}))^2,0)"
        )
    agreement["E10"] = "=SUM(B10:D10)"
    agreement.column_dimensions["A"].width = 64
    for column in "BCDE":
        agreement.column_dimensions[column].width = 24
    for row in range(3, 6):
        for column in range(2, 5):
            agreement.cell(row, column).number_format = "0.00%"
    agreement["B7"].number_format = "0.00%"
    agreement["B8"].number_format = "0.000"
    agreement["B9"].number_format = "0.000"
    agreement["E10"].number_format = "0.000"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--reuse-grid",
        action="store_true",
        help="Reuse a previously completed min_cluster_size grid and refit its winner only.",
    )
    args = parser.parse_args()

    import torch
    from sentence_transformers import SentenceTransformer

    OUTPUT.mkdir(parents=True, exist_ok=True)
    corpus = pd.read_csv(
        OUTPUT / "descriptive_corpus_sample.csv",
        usecols=["appid", "review", "voted_up"],
        dtype={"appid": "string", "review": "string"},
        keep_default_na=False,
    )
    corpus["voted_up"] = corpus["voted_up"].astype("int8")
    corpus = corpus.merge(load_catalog(), on="appid", how="left", validate="many_to_one")
    corpus["primary_genre"] = corpus["primary_genre"].fillna("Unknown")
    documents = corpus["review"].astype(str).tolist()
    device = "cuda" if torch.cuda.is_available() else "cpu"
    encoder = SentenceTransformer(EMBEDDING_MODEL, device=device)
    embeddings = encoder.encode(
        documents,
        batch_size=EMBEDDING_BATCH_SIZE,
        convert_to_numpy=True,
        show_progress_bar=True,
    )
    del encoder
    embeddings = np.asarray(embeddings, dtype=np.float32)

    grid_rows = []
    best_model = None
    best_words = None
    best_labels = None
    best_row = None
    started = time.perf_counter()
    grid_path = OUTPUT / "descriptive_min_cluster_grid.csv"
    if args.reuse_grid and grid_path.exists():
        existing_grid = pd.read_csv(grid_path)
        if sorted(existing_grid["min_cluster_size"].astype(int).tolist()) != sorted(
            MIN_CLUSTER_SIZES
        ):
            raise ValueError("The saved descriptive grid does not match the requested configurations.")
        grid_rows = existing_grid.to_dict(orient="records")
        eligible = existing_grid.loc[existing_grid["accepted"].astype(bool)]
        if not eligible.empty:
            best_row = eligible.sort_values("npmi", ascending=False).iloc[0].to_dict()
            config = {
                "seed": int(best_row["seed"]),
                "n_neighbors": int(best_row["n_neighbors"]),
                "n_components": int(best_row["n_components"]),
                "min_cluster_size": int(best_row["min_cluster_size"]),
                "min_samples": int(best_row["min_samples"]),
            }
            best_model = make_model(**config)
            best_labels, _ = best_model.fit_transform(documents, embeddings=embeddings)
            best_labels = np.asarray(best_labels, dtype=int)
            best_words = top_words_from_model(best_model)
    else:
        for min_cluster_size in MIN_CLUSTER_SIZES:
            config = {
                "seed": RANDOM_STATE,
                "n_neighbors": 15,
                "n_components": 5,
                "min_cluster_size": min_cluster_size,
                "min_samples": 5,
            }
            model = make_model(**config)
            labels, _ = model.fit_transform(documents, embeddings=embeddings)
            labels = np.asarray(labels, dtype=int)
            words = top_words_from_model(model)
            summary = summarize(documents, labels, words)
            accepted = (
                summary["topic_count"] >= 5
                and summary["largest_topic_share_non_outlier"] <= 0.50
            )
            row = {**config, **summary, "accepted": accepted}
            grid_rows.append(row)
            pd.DataFrame(grid_rows).to_csv(grid_path, index=False)
            if accepted and (
                best_row is None or summary["npmi"] > best_row["npmi"]
            ):
                best_model = model
                best_words = words
                best_labels = labels
                best_row = row
            if model is not best_model:
                del model

    if best_model is None or best_words is None or best_labels is None:
        result = {
            "status": "no_configuration_met_filters",
            "sample_rows": int(len(corpus)),
            "sample_games": int(corpus["appid"].nunique()),
            "tested_min_cluster_sizes": list(MIN_CLUSTER_SIZES),
            "grid": grid_rows,
        }
        (OUTPUT / "descriptive_topics_result.json").write_text(
            json.dumps(result, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return

    reduced_labels = np.asarray(
        best_model.reduce_outliers(
            documents,
            best_labels.tolist(),
            strategy="embeddings",
            embeddings=embeddings,
        ),
        dtype=int,
    )
    reassigned = int(np.sum((best_labels == -1) & (reduced_labels != -1)))
    raw_metrics = summarize(documents, best_labels, best_words)
    reduced_metrics = summarize(documents, reduced_labels, best_words)

    no_reduce = topic_genre_integration(corpus, best_labels, best_words)
    with_reduce = topic_genre_integration(corpus, reduced_labels, best_words)
    no_reduce.to_csv(OUTPUT / "integration_without_outlier_reduction.csv", index=False)
    with_reduce.to_csv(OUTPUT / "integration_with_outlier_reduction.csv", index=False)

    topic_rows = representative_topics(corpus, best_labels, embeddings, best_words)
    topic_rows.to_csv(OUTPUT / "descriptive_topics.csv", index=False)
    write_review_workbook(topic_rows, OUTPUT / "revision_temas.xlsx")
    for name, labels in (("without", best_labels), ("with", reduced_labels)):
        assignments = corpus[["appid", "voted_up", "primary_genre"]].copy()
        assignments["topic"] = labels
        assignments.to_csv(OUTPUT / f"descriptive_assignments_{name}_reduction.csv", index=False)

    result = {
        "status": "selected",
        "embedding_model": EMBEDDING_MODEL,
        "device": device,
        "sample_rows": int(len(corpus)),
        "sample_games": int(corpus["appid"].nunique()),
        "sample_class_counts": {
            str(int(key)): int(value)
            for key, value in corpus["voted_up"].value_counts().sort_index().items()
        },
        "selection_criterion": (
            "Highest NPMI among min_cluster_size values with K >= 5 "
            "and largest non-outlier topic <= 50%."
        ),
        "selected_min_cluster_size": int(best_row["min_cluster_size"]),
        "configuration": {
            "n_neighbors": 15,
            "n_components": 5,
            "min_cluster_size": int(best_row["min_cluster_size"]),
            "min_samples": 5,
            "metric_umap": "cosine",
            "metric_hdbscan": "euclidean",
            "random_state": RANDOM_STATE,
            "vectorizer_stop_words": "english",
            "vectorizer_ngram_range": [1, 1],
            "nr_topics": None,
        },
        "grid": grid_rows,
        "without_outlier_reduction": raw_metrics,
        "with_outlier_reduction": {
            **reduced_metrics,
            "reassigned_from_outlier": reassigned,
        },
        "elapsed_fit_seconds": float(time.perf_counter() - started),
    }
    (OUTPUT / "descriptive_topics_result.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
